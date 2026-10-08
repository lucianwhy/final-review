from contextlib import nullcontext
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from final_review.api import create_app
from final_review.material_jobs import process_material_job
from final_review.source_locators import located_sections


def _pdf_with_pages(texts: list[str]) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R 6 0 R] /Count 2 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, text in enumerate(texts):
        page_id = 4 + index * 2
        content_id = page_id + 1
        stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
        objects.extend(
            [
                (
                    f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
                ).encode(),
                b"<< /Length "
                + str(len(stream)).encode()
                + b" >>\nstream\n"
                + stream
                + b"\nendstream",
            ]
        )
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    start = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode()
    )
    return bytes(data)


def test_pdf_pages_and_ppt_slides_keep_real_ordinals():
    pages = located_sections(_pdf_with_pages(["First page", "Second page"]), "lesson.pdf")
    assert [(item["position"], item["text"].strip()) for item in pages] == [
        (1, "First page"),
        (2, "Second page"),
    ]

    presentation = Presentation()
    for text in ("First slide", "Second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text = text
    data = BytesIO()
    presentation.save(data)
    slides = located_sections(data.getvalue(), "lecture.pptx")
    assert [(item["position"], item["text"].strip()) for item in slides] == [
        (1, "First slide"),
        (2, "Second slide"),
    ]


def test_pdf_and_pptx_jobs_publish_positioned_chunks(system, tmp_path, monkeypatch):
    import pdfplumber

    from final_review import material_conversion

    system.settings.uploads_dir = str(tmp_path / "uploads")
    monkeypatch.setattr(
        material_conversion,
        "_convert_presentation_to_pdf",
        lambda _path, directory: directory / "material.pdf",
    )
    real_pdf_open = pdfplumber.open
    monkeypatch.setattr(
        pdfplumber,
        "open",
        lambda source: (
            nullcontext(SimpleNamespace(pages=[None, None]))
            if isinstance(source, Path)
            else real_pdf_open(source)
        ),
    )
    monkeypatch.setattr(material_conversion, "_find_executable", lambda *_args: "pdftoppm")
    picture = BytesIO()
    Image.new("RGB", (80, 40), "white").save(picture, format="PNG")
    monkeypatch.setattr(
        material_conversion,
        "_run",
        lambda command, *_args: (
            Path(command[-1]).with_suffix(".png").write_bytes(picture.getvalue())
        ),
    )
    monkeypatch.setattr(material_conversion, "_ocr_image", lambda *_args, **_kwargs: "")
    presentation = Presentation()
    for text in ("First slide", "Second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
    ppt = BytesIO()
    presentation.save(ppt)
    samples = [
        ("lecture.pdf", _pdf_with_pages(["First page", "Second page"]), "page"),
        ("lecture.pptx", ppt.getvalue(), "slide"),
    ]
    with TestClient(create_app(system.settings, system)) as client:
        for filename, content, kind in samples:
            uploaded = client.post(
                "/knowledge/upload",
                data={"course_id": "net", "title": filename, "source_type": "teacher_ppt"},
                files={"file": (filename, content)},
            ).json()
            job = system.store.claim_material_job()
            process_material_job(system.store, system.kb, job, 10 * 1024 * 1024)
            assert system.store.get_material_job(job["job_id"])["status"] == "succeeded"
            base = f"/api/courses/net/documents/{uploaded['document_id']}"
            listing = client.get(base + "/chunks").json()
            assert {(chunk["position_kind"], chunk["position"]) for chunk in listing["items"]} == {
                (kind, 1),
                (kind, 2),
            }


def test_chunk_preview_download_and_foreign_owner_are_guarded(system, tmp_path, monkeypatch):
    system.settings.uploads_dir = str(tmp_path / "uploads")
    with TestClient(create_app(system.settings, system)) as client:
        uploaded = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "Lecture", "source_type": "teacher_ppt"},
            files={"file": ("lesson.md", b"Source location matters")},
        ).json()
        job = system.store.claim_material_job()
        process_material_job(system.store, system.kb, job, 10 * 1024 * 1024)
        base = f"/api/courses/net/documents/{uploaded['document_id']}"

        def reject_backfill(document):
            raise AssertionError("preview must not write locators")

        monkeypatch.setattr(system.store, "ensure_material_source", reject_backfill)
        listing = client.get(base + "/chunks")
        assert listing.status_code == 200
        item = listing.json()["items"][0]
        assert item["position_kind"] == "document"
        assert item["text_start"] == 0
        assert client.get(base + "/chunks/unknown").status_code == 404
        detail = client.get(base + "/chunks/" + item["chunk_id"]).json()
        assert detail["content"] == "Source location matters"
        assert client.get(base + "/download").content == b"Source location matters"
        locator = system.store.get("source_locator", item["locator_id"])
        assert locator["material_version_id"] == listing.json()["material_version_id"]

        document = system.store.tables["document"][uploaded["document_id"]]
        document["user_id"] = "another-user"
        assert client.get(base + "/chunks").status_code == 404
        assert client.get(base + "/chunks/" + item["chunk_id"]).status_code == 404
        assert client.get(base + "/download").status_code == 404


def test_material_metadata_change_gets_new_version_without_moving_old_locator(system, tmp_path):
    system.settings.uploads_dir = str(tmp_path / "uploads")
    system.store.put(
        "course", "net", {"course_id": "net", "user_id": "local-user", "status": "active"}
    )
    with TestClient(create_app(system.settings, system)) as client:
        uploaded = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "Original", "source_type": "homework"},
            files={"file": ("lesson.md", b"A stable excerpt")},
        ).json()
        job = system.store.claim_material_job()
        process_material_job(system.store, system.kb, job, 10 * 1024 * 1024)
        base = f"/api/courses/net/documents/{uploaded['document_id']}"
        before = client.get(base + "/chunks").json()
        document = system.store.get("document", uploaded["document_id"])
        changed = client.patch(
            base,
            json={
                "title": "Renamed",
                "chapter": "Chapter 2",
                "source_type": "teacher_ppt",
                "expected_updated_at": document["updated_at"],
            },
        )
        assert changed.status_code == 200
        after = client.get(base + "/chunks").json()
        assert after["material_version_id"] != before["material_version_id"]
        assert after["source_type"] == "teacher_ppt"
        old_version = system.store.get("material_version", before["material_version_id"])
        assert old_version["source_type"] == "homework"


def test_existing_text_ingest_has_a_previewable_document_locator(system):
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/ingest",
            json={
                "course_id": "net",
                "title": "Typed notes",
                "source_type": "homework",
                "markdown": "One useful paragraph",
            },
        )
        assert response.status_code == 200
        document_id = response.json()["document_id"]
        preview = client.get(f"/api/courses/net/documents/{document_id}/chunks")
        assert preview.status_code == 200
        assert preview.json()["items"][0]["position_kind"] == "document"
