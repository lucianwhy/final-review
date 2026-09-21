from fastapi.testclient import TestClient

from final_review.api import create_app


def test_api_full_feedback_cycle(system):
    with TestClient(create_app(system.settings, system)) as client:
        assert client.get("/health").status_code == 200
        assert client.get("/openapi.json").status_code == 200
        response = client.post(
            "/agent/invoke",
            json={
                "course_id": "net",
                "session_id": "http",
                "message": "出题",
                "intent": "quiz",
            },
        )
        assert response.json()["status"] == "needs_input"
        result = client.post(
            "/agent/resume",
            json={
                "course_id": "net",
                "session_id": "http",
                "exam_profile": {"question_types": ["short_answer"]},
            },
        ).json()
        assert result["status"] == "awaiting_answers"
        assert "reference_answer" not in str(result)
        graded = client.post(
            "/assessment/evaluate",
            json={
                "course_id": "net",
                "session_id": "http",
                "answers": {q["id"]: "不清楚" for q in result["questions"]},
            },
        )
        assert graded.status_code == 200
        assert graded.json()["assessment"]["score"] == 30


def test_upload_supported_and_reject_bad_file(system):
    with TestClient(create_app(system.settings, system)) as client:
        data = {"course_id": "net", "title": "上传", "source_type": "homework"}
        assert (
            client.post(
                "/knowledge/upload", data=data, files={"file": ("note.md", "测试内容".encode())}
            ).status_code
            == 200
        )
        assert (
            client.post(
                "/knowledge/upload", data=data, files={"file": ("test.exe", b"MZ")}
            ).status_code
            == 422
        )
        assert (
            client.post(
                "/knowledge/upload", data=data, files={"file": ("empty.md", b"")}
            ).status_code
            == 422
        )


def test_auth_and_missing_session(system):
    from pydantic import SecretStr

    system.settings.api_token = SecretStr("test-token")
    with TestClient(create_app(system.settings, system)) as client:
        url = "/agent/recover?course_id=net&session_id=missing"
        assert client.post(url).status_code == 401
        assert client.post(url, headers={"Authorization": "Bearer test-token"}).status_code == 404


def test_invalid_request_never_reaches_model(system):
    with TestClient(create_app(system.settings, system)) as client:
        result = client.post(
            "/agent/invoke",
            json={
                "course_id": "net",
                "session_id": "test",
                "message": "出题",
                "intent": "unknown",
            },
        )
        assert result.status_code == 422
        assert system.model.retrieval_calls == 0


def test_chat_reports_missing_deepseek_key(system):
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post("/api/chat", json={"message": "帮我复习需求分析"})
        assert response.status_code == 503
        assert response.json()["detail"] == "尚未配置 DEEPSEEK_API_KEY"


def test_real_pptx_conversion(system):
    from io import BytesIO

    from pptx import Presentation

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "TCP 三次握手"
    slide.placeholders[1].text = "同步双方初始序列号并确认双方收发能力。"
    data = BytesIO()
    slides.save(data)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={
                "course_id": "net",
                "title": "转换测试",
                "source_type": "teacher_ppt",
            },
            files={"file": ("slides.pptx", data.getvalue())},
        )
        assert response.status_code == 200, response.text
        stored = system.store.get("document", response.json()["document_id"])
        assert "三次握手" in stored["cleaned_markdown"]
