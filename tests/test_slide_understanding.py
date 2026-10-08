from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from final_review.api import create_app
from final_review.material_conversion import ConvertedMaterial
from final_review.slide_understanding import SlideInterpreter


def interpreter(replies):
    instance = SlideInterpreter.__new__(SlideInterpreter)
    instance.identity = "test-model-and-prompt"
    instance.model = "test-vision"
    instance.repairs = 1
    instance.calls = []

    def ask(prompt, image, text):
        instance.calls.append((prompt, image, text))
        return next(replies)

    instance._ask = ask
    return instance


def reading(**changes):
    return {
        "kind": "knowledge",
        "title": "请求处理",
        "blocks": [
            {
                "title": "请求对象",
                "text": "容器创建请求对象，并调用 service() 方法。",
                "evidence": [{"source_id": "image-page", "quote": "容器创建请求对象"}],
            }
        ],
        "uncertainties": [],
        **changes,
    }


def review(**changes):
    issues = changes.get("issues", [])
    return {
        "block_reviews": changes.get(
            "block_reviews",
            [
                {
                    "block_id": "b1",
                    "supported": changes.get("faithful", True),
                    "readable": changes.get("readable", True),
                    "uncertainties_resolved": changes.get("uncertainties_resolved", True),
                    "issues": [
                        {"kind": "unsupported_fact", "message": str(issue)} for issue in issues
                    ],
                }
            ],
        ),
        "missing_evidence_ids": changes.get("missing_evidence_ids", []),
    }


def test_verified_page_cache_avoids_repeated_model_calls(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image bytes")
    reader = interpreter(iter([reading(), review()]))
    result = reader.read(image, "原生文字", "备注", tmp_path / "cache")
    assert result["quality"] == "verified"
    assert len(reader.calls) == 2
    assert reader.calls[0][1] == b"image bytes"
    assert reader.read(image, "原生文字", "备注", tmp_path / "cache") == result
    assert len(reader.calls) == 2
    reader.identity = "changed-model"
    with pytest.raises(StopIteration):
        reader.read(image, "原生文字", "备注", tmp_path / "cache")


def test_unfaithful_page_repairs_then_remains_excluded(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(
        iter(
            [
                reading(),
                review(faithful=False, issues=["图中没有此步骤"]),
                reading(
                    blocks=[
                        {
                            "title": "请求对象",
                            "text": "容器创建请求对象",
                            "evidence": [{"source_id": "image-page", "quote": "容器创建请求对象"}],
                            "uncertainties": ["箭头不清楚"],
                        }
                    ]
                ),
                review(uncertainties_resolved=False),
            ]
        )
    )
    result = reader.read(image, "文字", "", tmp_path / "cache")
    assert result["quality"] == "review_needed"
    assert "箭头不清楚" in result["blocks"][0]["issues"]
    assert "图中没有此步骤" in reader.calls[2][2]


def test_corrupt_model_output_is_not_accepted_even_if_reviewer_passes(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    broken = reading(
        blocks=[
            {
                "title": "流程",
                "markdown": "乱码\ufffd service()()",
                "evidence": [{"source_id": "image-page", "quote": "service()"}],
            }
        ]
    )
    reader = interpreter(iter([broken, review(), broken, review()]))
    assert reader.read(image, "", "", tmp_path / "cache")["quality"] == "review_needed"


def test_navigation_page_is_verified_without_knowledge_blocks(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(
        iter([reading(kind="navigation", blocks=[]), {"is_navigation": True, "issues": []}])
    )
    result = reader.read(image, "目录", "", tmp_path / "cache")
    assert result["quality"] == "verified"
    assert result["blocks"] == []


def test_false_navigation_classification_cannot_hide_teaching_content(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(
        iter(
            [
                reading(kind="navigation", blocks=[]),
                {"is_navigation": False, "issues": ["原页包含定义"]},
            ]
        )
    )
    reader.repairs = 0
    result = reader.read(image, "知识定义", "", tmp_path / "cache")
    assert result["quality"] == "review_needed"
    assert result["issues"] == ["原页包含定义"]


def test_final_plain_text_keeps_code_symbols_and_removes_prose_markup():
    from final_review.plain_material_text import plain_material_text

    source = "# 标题\n\n**定义**与 `service()`\n```css\n#main { color: #fff; }\n```\n- 普通条目"
    assert (
        plain_material_text(source) == "标题\n\n定义与 service()\n#main { color: #fff; }\n普通条目"
    )


def test_understood_plain_code_is_not_stripped_again(system):
    from final_review.schemas import MaterialInput

    source = "示例代码\n# Python comment\nx = 2 ** 3\n#main { color: #fff; }"
    document, chunks = system.kb.prepare(
        MaterialInput(
            course_id="net",
            title="代码页",
            source_type="teacher_ppt",
            markdown=source,
        ),
        sections=[{"text": source, "position_kind": "slide", "position": 1}],
        plain_text=True,
    )
    assert chunks[0]["content"] == source
    assert document["cleaned_markdown"] == source


def test_literal_code_symbols_are_preserved_but_added_markup_is_rejected(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    code = "# Python comment\nx = 2 ** 3"
    candidate = reading(
        blocks=[
            {
                "title": "原生代码",
                "text": code,
                "evidence": [
                    {"source_id": "native-1", "quote": "# Python comment"},
                    {"source_id": "native-2", "quote": "x = 2 ** 3"},
                ],
            }
        ]
    )
    result = interpreter(iter([candidate, review()])).read(image, code, "", tmp_path / "cache")
    assert result["quality"] == "verified"
    assert result["blocks"][0]["text"] == code
    candidate = reading(
        blocks=[
            {
                "title": "擅加格式",
                "text": "**定义**",
                "evidence": [{"source_id": "native-1", "quote": "定义"}],
            }
        ]
    )
    reader = interpreter(iter([candidate, review()]))
    reader.repairs = 0
    assert reader.read(image, "定义", "", tmp_path / "markup")["quality"] == "review_needed"


def test_malformed_model_reply_retries_without_inserting_raw_text(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(iter([{}, reading(), review(issues=[{"reason": "不能核对"}])]))
    result = reader.read(image, "原始乱码", "", tmp_path / "cache")
    assert result["quality"] == "review_needed"
    assert result["raw_native"] == "原始乱码"


def test_evidence_separates_native_notes_and_image_ocr():
    from final_review.slide_understanding import source_evidence

    sources = source_evidence(
        "Servlet\nservice()", "Servlet\n### Notes:\n教师强调大小写", "GET读取资源"
    )
    assert sources == [
        {"source_id": "native-1", "origin": "native", "text": "Servlet"},
        {"source_id": "native-2", "origin": "native", "text": "service()"},
        {"source_id": "note-1", "origin": "note", "text": "教师强调大小写"},
        {
            "source_id": "image-page",
            "origin": "image",
            "text": "原页图片见附图",
            "ocr_hint": "GET读取资源",
            "region": [0, 0, 1000, 1000],
        },
    ]


def test_partial_page_retains_good_block_and_rejects_fabricated_fact(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    candidate = reading(
        blocks=[
            {
                "title": "正确代码",
                "text": 'request.getParameter("name")',
                "evidence": [{"source_id": "native-1", "quote": 'request.getParameter("name")'}],
            },
            {
                "title": "伪造结论",
                "text": "GET加密所有请求",
                "evidence": [{"source_id": "image-page", "quote": "GET用于读取资源"}],
            },
        ]
    )
    verdict = review(
        block_reviews=[
            {"block_id": "b1", "supported": True, "readable": True, "issues": []},
            {
                "block_id": "b2",
                "supported": False,
                "readable": True,
                "issues": [{"kind": "unsupported_fact", "message": "图片并未说明加密"}],
            },
        ]
    )
    reader = interpreter(iter([candidate, verdict, {}]))
    result = reader.read(image, 'request.getParameter("name")', "", tmp_path / "cache")
    assert result["quality"] == "partial"
    assert [block["quality"] for block in result["blocks"]] == ["verified", "review_needed"]
    assert result["blocks"][0]["evidence"][0]["origin"] == "native"


@pytest.mark.parametrize(
    "reference",
    [
        {"source_id": "missing-source", "quote": "原文"},
        {"source_id": "native-1", "quote": "不存在的原文"},
    ],
)
def test_unknown_or_forged_native_quote_rejected_even_if_model_passes(tmp_path, reference):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    candidate = reading(blocks=[{"title": "知识", "text": "原文", "evidence": [reference]}])
    reader = interpreter(iter([candidate, review()]))
    reader.repairs = 0
    assert reader.read(image, "原文", "", tmp_path / "cache")["quality"] == "review_needed"


def test_image_only_quote_is_not_required_in_native_text_and_layout_is_warning(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    candidate = reading(
        blocks=[
            {
                "title": "HTTP方法",
                "text": "GET用于读取资源",
                "evidence": [{"source_id": "image-page", "quote": "GET用于读取资源"}],
            }
        ]
    )
    verdict = review(
        block_reviews=[
            {
                "block_id": "b1",
                "supported": True,
                "readable": True,
                "issues": [{"kind": "layout", "message": "图片文字颜色不同"}],
            }
        ]
    )
    result = interpreter(iter([candidate, verdict])).read(
        image, "Servlet代码", "", tmp_path / "cache"
    )
    assert result["quality"] == "verified"
    assert result["issues"] == []
    assert result["blocks"][0]["warnings"] == ["图片文字颜色不同"]


@pytest.mark.parametrize("kind", ["code_mismatch", "unreadable"])
def test_bad_code_or_unreadable_block_cannot_pass(tmp_path, kind):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    verdict = review(
        block_reviews=[
            {
                "block_id": "b1",
                "supported": True,
                "readable": True,
                "issues": [{"kind": kind, "message": "无法核对"}],
            }
        ]
    )
    reader = interpreter(iter([reading(), verdict]))
    reader.repairs = 0
    assert reader.read(image, "", "", tmp_path / "cache")["quality"] == "review_needed"


@pytest.mark.parametrize("ids", [[], ["b1", "b1"], ["unknown"]])
def test_incomplete_duplicate_or_unknown_block_verdicts_fail_closed(tmp_path, ids):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    verdict = review(
        block_reviews=[
            {"block_id": key, "supported": True, "readable": True, "issues": []} for key in ids
        ]
    )
    reader = interpreter(iter([reading(), verdict]))
    reader.repairs = 0
    assert reader.read(image, "", "", tmp_path / "cache")["quality"] == "review_needed"


def test_missing_content_is_visible_without_discarding_verified_blocks(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    result = interpreter(iter([reading(), review(missing_evidence_ids=["native-1"]), {}])).read(
        image, "尚未整理的重要内容", "", tmp_path / "cache"
    )
    assert result["quality"] == "partial"
    assert result["blocks"][0]["quality"] == "verified"
    assert "遗漏教学证据 native-1" in result["issues"]


def test_reviewer_must_explicitly_resolve_reader_uncertainties(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    candidate = reading(
        uncertainties=["OCR 提示有误，需要对照原图"],
        blocks=[
            {
                "title": "代码",
                "text": "service()",
                "uncertainties": ["OCR 提示重复括号"],
                "evidence": [{"source_id": "native-1", "quote": "service()"}],
            }
        ],
    )
    verdict = review(
        block_reviews=[
            {
                "block_id": "b1",
                "supported": True,
                "readable": True,
                "uncertainties_resolved": True,
                "issues": [],
            }
        ]
    )
    verdict["page_uncertainties_resolved"] = True
    result = interpreter(iter([candidate, verdict])).read(
        image, "service()", "", tmp_path / "cache"
    )
    assert result["quality"] == "verified"
    assert result["blocks"][0]["uncertainties"] == ["OCR 提示重复括号"]
    assert result["review"]["page_uncertainties_resolved"] is True


def test_image_cache_tracks_each_pages_source_even_after_partial_rebuild(tmp_path, monkeypatch):
    from final_review import material_conversion

    calls = []

    def render(command, *_args):
        calls.append(command)
        Path(command[-1]).with_suffix(".png").write_bytes(f"render-{len(calls)}".encode())

    monkeypatch.setattr(material_conversion, "_run", render)

    class Reader:
        def read(self, *_args):
            return {
                "title": "知识",
                "kind": "knowledge",
                "quality": "verified",
                "issues": [],
                "blocks": [{"title": "内容", "text": "实际知识", "quality": "verified"}],
            }

    def convert(signature, pages=None):
        return material_conversion._understand_rendered_slides(
            tmp_path / "deck.pdf",
            "renderer",
            [{"text": "相同原生标题"}] * 2,
            [""] * 2,
            Reader(),
            tmp_path / "cache",
            None,
            pages,
            source_signature=signature,
        )

    convert("original")
    convert("original")
    assert len(calls) == 2
    convert("image-only-change", {1})
    assert len(calls) == 3
    convert("image-only-change")
    assert len(calls) == 4
    assert (tmp_path / "cache/slide-002.source").read_text() == "image-only-change"


def test_visual_conversion_excludes_unreviewed_and_navigation_pages(tmp_path, monkeypatch):
    from io import BytesIO

    import pdfplumber
    from pptx import Presentation
    from pptx.util import Inches

    from final_review import material_conversion

    ppt = Presentation()
    for text in ["目录", "正常定义", "错字乱码", "部分通过"]:
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = text
    buffer = BytesIO()
    ppt.save(buffer)
    monkeypatch.setattr(
        material_conversion,
        "_convert_presentation_to_pdf",
        lambda _path, folder: folder / "test.pdf",
    )
    from contextlib import nullcontext

    monkeypatch.setattr(
        pdfplumber, "open", lambda _path: nullcontext(SimpleNamespace(pages=[None] * 4))
    )
    monkeypatch.setattr(material_conversion, "_find_executable", lambda *_args: "pdftoppm")
    monkeypatch.setattr(
        material_conversion,
        "_run",
        lambda cmd, *_args: Path(cmd[-1]).with_suffix(".png").write_bytes(b"image"),
    )

    class Reader:
        def read(self, image, native, extracted, cache):
            if native == "部分通过":
                return {
                    "title": native,
                    "kind": "knowledge",
                    "quality": "partial",
                    "issues": ["错误块不能入库"],
                    "blocks": [
                        {
                            "title": "局部正确内容",
                            "text": "保留真实代码",
                            "quality": "verified",
                            "block_id": "b1",
                            "issues": [],
                        },
                        {
                            "title": "伪造内容",
                            "text": "不应被检索的假事实",
                            "quality": "review_needed",
                            "block_id": "b2",
                            "issues": ["来源不符"],
                        },
                    ],
                }
            return {
                "title": native,
                "kind": "navigation" if native == "目录" else "knowledge",
                "quality": "review_needed" if native == "错字乱码" else "verified",
                "issues": [],
                "blocks": [{"title": "整理后知识点", "markdown": "完整定义"}],
            }

    result = material_conversion.convert_material(
        buffer.getvalue(), "slides.pptx", 10**7, interpreter=Reader(), cache_dir=tmp_path / "pages"
    )
    assert len(result.sections) == 2
    assert result.sections[0]["position"] == 2
    assert "完整定义" in result.markdown
    assert "错字乱码" not in result.markdown
    assert "目录" not in result.markdown
    assert "保留真实代码" in result.markdown
    assert "不应被检索的假事实" not in result.markdown
    assert len(result.pages) == 4
    assert result.pages[3]["quality"] == "partial"
    assert result.pages[3]["blocks"][1]["issues"] == ["来源不符"]


@pytest.mark.parametrize("quality", ["verified", "partial"])
def test_preview_exposes_quality_and_protects_original_page_path(
    system, tmp_path, monkeypatch, quality
):
    from final_review import material_jobs

    system.settings.uploads_dir = str(tmp_path)
    monkeypatch.setattr(
        material_jobs,
        "convert_material",
        lambda *_args, **_kwargs: ConvertedMaterial(
            markdown="整理后的真实知识",
            sections=[{"text": "整理后的真实知识", "position_kind": "slide", "position": 1}],
            pages=[
                {
                    "position": 1,
                    "title": "知识页",
                    "kind": "knowledge",
                    "quality": quality,
                    "issues": [],
                }
            ],
            pipeline="visual-slides-v1",
        ),
    )
    with TestClient(create_app(system.settings, system)) as client:
        uploaded = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "课件", "source_type": "teacher_ppt"},
            files={"file": ("slides.pptx", b"test")},
        )
        job = system.store.claim_material_job()
        material_jobs.process_material_job(system.store, system.kb, job, 10**7)
        base = f"/api/courses/net/documents/{uploaded.json()['document_id']}"
        preview = client.get(base + "/chunks").json()
        assert preview["quality_status"] == quality
        document = system.store.get("document", uploaded.json()["document_id"])
        directory = Path(document["file_path"]).with_suffix(".analysis")
        directory.mkdir()
        (directory / "slide-001.png").write_bytes(b"PNG")
        assert client.get(base + "/pages/1").content == b"PNG"
        assert client.get(base + "/pages/2").status_code == 404
        assert client.get(base.replace("net", "other") + "/pages/1").status_code == 404


def test_billing_failure_is_actionable_and_never_indexes_raw_ppt(system, tmp_path, monkeypatch):
    from final_review import material_jobs
    from final_review.slide_understanding import VisionServiceUnavailable

    system.settings.uploads_dir = str(tmp_path)
    system.settings.material_vision_enabled = True

    def unavailable(_settings):
        raise VisionServiceUnavailable("视觉模型账号欠费，请恢复后重试")

    monkeypatch.setattr(material_jobs, "SlideInterpreter", unavailable)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={
                "course_id": "net",
                "title": "课件",
                "source_type": "teacher_ppt",
            },
            files={"file": ("slides.pptx", b"raw ppt")},
        )
        job = system.store.claim_material_job()
        material_jobs.process_material_job(system.store, system.kb, job, 10**7)
        result = system.store.get_material_job(job["job_id"])
        assert result["status"] == "failed"
        assert result["error_code"] == "vision_unavailable"
        assert "欠费" in result["error_message"]
        assert not any(
            chunk["document_id"] == response.json()["document_id"] for chunk in system.store.chunks
        )
