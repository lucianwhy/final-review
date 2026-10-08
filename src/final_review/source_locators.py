"""Stable source positions for searchable material chunks."""

from hashlib import sha256
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from .storage import stable_key


def located_sections(content: bytes, filename: str) -> list[dict] | None:
    """Return page/slide text only when the original position is reliable."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        import pdfplumber

        try:
            with pdfplumber.open(BytesIO(content)) as pdf:
                sections = [
                    {"position_kind": "page", "position": number, "text": page.extract_text() or ""}
                    for number, page in enumerate(pdf.pages, 1)
                ]
        except Exception:
            return None
    elif suffix in {".ppt", ".pptx"}:
        from pptx import Presentation

        try:
            if suffix == ".ppt":
                from .material_conversion import _convert_legacy

                with TemporaryDirectory(prefix="final-review-locator-") as temporary:
                    directory = Path(temporary)
                    source = directory / "material.ppt"
                    source.write_bytes(content)
                    presentation = Presentation(_convert_legacy(source, directory))
                    sections = _slide_sections(presentation)
            else:
                sections = _slide_sections(Presentation(BytesIO(content)))
        except Exception:
            return None
    else:
        return None
    return sections if any(section["text"].strip() for section in sections) else None


def _slide_sections(presentation) -> list[dict]:
    sections = []
    for number, slide in enumerate(presentation.slides, 1):
        text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
            if shape.has_table:
                text.extend(cell.text for row in shape.table.rows for cell in row.cells)
        sections.append(
            {
                "position_kind": "slide",
                "position": number,
                "text": "\n".join(part for part in text if part.strip()),
            }
        )
    return sections


def material_version(document: dict) -> dict:
    version_id = document.get("material_version_id") or stable_key(
        document["document_id"],
        document.get("cleaned_markdown", ""),
        document.get("title", ""),
        document.get("chapter", ""),
        document.get("source_type", ""),
    )
    return {
        "material_version_id": version_id,
        "document_id": document["document_id"],
        "course_id": document["course_id"],
        "user_id": document["user_id"],
        "file_name": document.get("file_name") or document["title"],
        "source_type": document["source_type"],
        "source_origin": document.get("source_origin", "legacy_upload"),
        "content_hash": sha256(document.get("cleaned_markdown", "").encode()).hexdigest(),
        "created_at": document.get("uploaded_at") or document.get("created_at"),
    }


def chunk_locator(version: dict, chunk: dict, ordinal: int) -> dict:
    return {
        "user_id": version["user_id"],
        "course_id": version["course_id"],
        "material_version_id": version["material_version_id"],
        "locator_id": chunk["chunk_id"],
        "locator_kind": "chunk",
        "ordinal": ordinal,
        "data": {
            "chunk_id": chunk["chunk_id"],
            "position_kind": chunk.get("position_kind", "document"),
            "position": chunk.get("position"),
            "text_start": chunk.get("text_start"),
            "text_end": chunk.get("text_end"),
            "content": chunk.get("content", ""),
        },
    }
