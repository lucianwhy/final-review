"""M1-09 acceptance: real vision/embedding, isolated PostgreSQL, failure injection.

A randomly created database, evidence directory and test-named CLI logs are written.
Run from the project root with uv run python scripts/check_m1_09_acceptance.py.
"""

import argparse
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import time
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import psycopg
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from psycopg import sql
from psycopg.conninfo import make_conninfo
from pydantic import SecretStr

from final_review.api import create_app
from final_review.config import Settings
from final_review.migrations import apply_migrations
from final_review.postgres import PostgresStore
from final_review.rag import KnowledgeBase
from final_review.schemas import MaterialInput

ROOT = Path(__file__).resolve().parents[1]


def replay_main():
    """Replay already-recorded real conversion to isolate database checks."""
    from final_review.material_conversion import ConvertedMaterial
    from scripts import rebuild_materials

    candidate = Path(os.environ["M109_REPLAY_JSON"])
    data = json.loads(candidate.read_text(encoding="utf-8"))
    source = Path(os.environ["M109_REPLAY_PPTX"]).read_bytes()
    original = rebuild_materials.rebuild

    def replay(document, settings, *, pages=None):
        if Path(document["file_path"]).read_bytes() != source:
            return original(document, settings, pages=pages)
        sections = [item for item in data["sections"] if pages is None or item["position"] in pages]
        page_rows = [item for item in data["pages"] if pages is None or item["position"] in pages]
        converted = ConvertedMaterial(
            markdown="\n\n".join(item["text"] for item in sections),
            sections=sections,
            pages=page_rows,
            pipeline="visual-slides-v1",
        )
        report = Path(document["file_path"]).with_suffix(".analysis") / "candidate.json"
        report.write_text(
            json.dumps(
                {"markdown": converted.markdown, "sections": sections, "pages": page_rows},
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return converted, report

    with patch.object(rebuild_materials, "rebuild", side_effect=replay):
        rebuild_materials.main()


def replay_course_main():
    from scripts import rebuild_course_materials

    run = subprocess.run

    def child(command, **kwargs):
        command = [
            command[0],
            "-c",
            "from scripts.check_m1_09_acceptance import replay_main; replay_main()",
            *command[2:],
        ]
        return run(command, **kwargs)

    with patch.object(rebuild_course_materials.subprocess, "run", side_effect=child):
        rebuild_course_materials.main()


class FixedEmbeddings:
    def embed_documents(self, texts):
        return [[1.0] + [0.0] * 1535 for _ in texts]


def make_deck(directory):
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 44)
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    for number, text in enumerate(
        [
            "HTTP 会话：图片知识页",
            "Cookie 保存在客户端。",
            "Session 状态保存在服务端。",
            "HTTP 是无状态协议。",
        ]
    ):
        draw.text((90, 100 + number * 145), text, font=font, fill="#18322a")
    picture = directory / "image-only.png"
    image.save(picture)
    image = Image.new("RGB", (1400, 300), "white")
    ImageDraw.Draw(image).text(
        (40, 80), "GET 用于读取资源，POST 用于提交数据。", font=font, fill="#18322a"
    )
    mixed = directory / "mixed-image.png"
    image.save(mixed)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    navigation = deck.slides.add_slide(deck.slide_layouts[6])
    navigation.shapes.add_textbox(
        Inches(1), Inches(1), Inches(11), Inches(4)
    ).text = "目录\n1. HTTP 会话\n2. Servlet 请求处理"
    pure = deck.slides.add_slide(deck.slide_layouts[6])
    pure.shapes.add_picture(str(picture), 0, 0, width=deck.slide_width, height=deck.slide_height)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(4)).text_frame
    frame.text = (
        'Servlet 请求处理\nrequest.getParameter("name") 读取参数。\n'
        'System.out.println("ok");\n响应状态码 404 表示资源未找到。'
    )
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(26)
    slide.shapes.add_picture(str(mixed), Inches(0.6), Inches(4.7), width=Inches(12))
    slide.notes_slide.notes_text_frame.text = "教师备注：参数名 name 区分大小写。"
    path = directory / "controlled-slides.pptx"
    deck.save(path)
    return path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--reuse-fixture", type=Path, help="Prior test evidence directory")
    parser.add_argument(
        "--replay-candidate",
        action="store_true",
        help="Replay prior real conversion for independent DB checks",
    )
    arguments = parser.parse_args()
    if arguments.replay_candidate and not arguments.reuse_fixture:
        parser.error("--replay-candidate requires --reuse-fixture")
    settings = Settings()
    root_url = settings.database_url.get_secret_value()
    database = "final_review_m109_" + uuid4().hex[:12]
    isolated_url = make_conninfo(root_url, dbname=database)
    directory = ROOT / "backups" / ("m1-09-acceptance-" + uuid4().hex[:8])
    directory.mkdir(parents=True)
    report = {
        "database": "isolated disposable PostgreSQL",
        "checks": [],
        "vision_model": settings.material_vision_model,
        "fixtures": "controlled real PPTX",
    }
    environment = {**os.environ, "DATABASE_URL": isolated_url, "PYTHONIOENCODING": "utf-8"}
    report["conversion_replayed"] = arguments.replay_candidate
    if arguments.replay_candidate:
        previous = arguments.reuse_fixture.resolve()
        environment.update(
            M109_REPLAY_JSON=str(previous / "controlled-slides.analysis/candidate.json"),
            M109_REPLAY_PPTX=str(previous / "controlled-slides.pptx"),
        )
    started = time.monotonic()
    store = None

    def check(name, condition, *, visual=False):
        if visual and not condition:
            report.setdefault("failed_visual_checks", []).append(name)
            print("FAIL", name, flush=True)
            return
        assert condition, name
        report["checks"].append(name)
        print("PASS", name, flush=True)

    def cli(arguments, expected=0):
        command = [sys.executable, *arguments]
        if report["conversion_replayed"]:
            function = "replay_course_main" if "rebuild_course" in arguments[0] else "replay_main"
            command = [
                sys.executable,
                "-c",
                f"from scripts.check_m1_09_acceptance import {function}; {function}()",
                *arguments[1:],
            ]
        process = subprocess.run(
            command,
            cwd=ROOT,
            env=environment,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=600,
            check=False,
        )
        # Save diagnostics from test-only data, never the connection settings.
        (directory / f"cli-{len(list(directory.glob('cli-*.log'))):02d}.log").write_text(
            process.stdout + process.stderr, encoding="utf-8"
        )
        check("CLI " + " ".join(arguments) + f" exit={expected}", process.returncode == expected)
        return process

    with psycopg.connect(root_url, autocommit=True) as admin:
        admin.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(database)))
    try:
        apply_migrations(isolated_url)
        api_settings = settings.model_copy(
            update={
                "database_url": SecretStr(isolated_url),
                "uploads_dir": str(directory),
                "auth_cookie_secure": False,
                "llm_api_key": SecretStr(""),
                "embedding_api_key": SecretStr(""),
            }
        )
        app = create_app(api_settings)
        with TestClient(app) as owner:
            other = TestClient(app)
            owner_result = owner.post(
                "/api/auth/sign-up",
                json={"email": "owner@m109.test", "password": "m109-test-password"},
            )
            check("owner registration", owner_result.status_code == 200)
            check(
                "second user registration",
                other.post(
                    "/api/auth/sign-up",
                    json={"email": "other@m109.test", "password": "m109-test-password"},
                ).status_code
                == 200,
            )
            course = owner.post("/api/courses", json={"name": "M1-09 隔离验收"}).json()
            second_course = owner.post("/api/courses", json={"name": "其他课程"}).json()[
                "course_id"
            ]
            store = PostgresStore(isolated_url)
            store.setup()
            store.bind_user(course["user_id"])
            kb = KnowledgeBase(store, FixedEmbeddings(), settings)
            path = make_deck(directory)
            if arguments.reuse_fixture:
                previous = arguments.reuse_fixture.resolve()
                shutil.copy2(previous / path.name, path)
                shutil.copytree(
                    previous / path.with_suffix(".analysis").name, path.with_suffix(".analysis")
                )

            def seed(document_id, file_path):
                document, chunks = kb.prepare(
                    MaterialInput(
                        document_id=document_id,
                        course_id=course["course_id"],
                        title=document_id,
                        source_type="teacher_ppt",
                        markdown="LEGACY_EXCERPT_SENTINEL：旧版引用内容",
                    )
                )
                document.update(
                    file_name=file_path.name,
                    file_path=str(file_path),
                    parse_status="ready",
                    user_id=course["user_id"],
                )
                store.ingest(document, chunks)
                return store.get("document", document_id), store.list_material_chunks(document_id)

            old_document, old_chunks = seed("m109-controlled", path)
            cli(
                [
                    "scripts/rebuild_materials.py",
                    "--document-id",
                    "m109-controlled",
                    "--pages",
                    "2",
                    "--publish",
                ],
                expected=2,
            )
            check(
                "partial publish leaves document unchanged",
                store.get("document", "m109-controlled") == old_document,
            )
            cli(
                ["scripts/rebuild_materials.py", "--document-id", "m109-controlled", "--pages", "2"]
            )
            check(
                "preview candidate keeps old index",
                store.list_material_chunks("m109-controlled") == old_chunks,
            )
            cli(["scripts/rebuild_materials.py", "--document-id", "m109-controlled", "--publish"])
            current = store.get("document", "m109-controlled")
            chunks = store.list_material_chunks("m109-controlled")
            report["pages"] = current["analysis_pages"]
            report["cleaned_text"] = current["cleaned_markdown"]
            check(
                "navigation page excluded; image and mixed pages indexed",
                {item["position"] for item in chunks} == {2, 3}
                and current["analysis_pages"][0]["kind"] == "navigation",
                visual=True,
            )
            text = current["cleaned_markdown"]
            check(
                "image facts, native code and image facts retained",
                all(
                    term in text
                    for term in [
                        "Cookie",
                        "Session",
                        "getParameter",
                        '"name"',
                        "System.out.println",
                        '"ok"',
                        "404",
                        "GET",
                        "POST",
                    ]
                ),
                visual=True,
            )
            check(
                "all controlled pages verified",
                all(page["quality"] == "verified" for page in current["analysis_pages"]),
                visual=True,
            )
            check(
                "historical reference preserved outside active index",
                old_chunks[0]["chunk_id"]
                in {item["chunk_id"] for item in current["historical_chunks"]}
                and not any("LEGACY_EXCERPT_SENTINEL" in item["content"] for item in chunks),
            )
            base = f"/api/courses/{course['course_id']}/documents/m109-controlled"
            for page in (1, 2, 3):
                response = owner.get(base + f"/pages/{page}")
                check(
                    f"owner original page {page}",
                    response.status_code == 200 and response.content.startswith(b"\x89PNG"),
                )
                check(
                    f"other user page {page} denied",
                    other.get(base + f"/pages/{page}").status_code == 404,
                )
            check(
                "cross course original denied",
                owner.get(base.replace(course["course_id"], second_course) + "/pages/2").status_code
                == 404,
            )
            check("missing original page denied", owner.get(base + "/pages/999").status_code == 404)
            check(
                "owner history reference readable",
                owner.get(base + "/chunks/" + old_chunks[0]["chunk_id"]).json().get("historical")
                is True,
            )
            check(
                "other user history denied",
                other.get(base + "/chunks/" + old_chunks[0]["chunk_id"]).status_code == 404,
            )
            check(
                "owner preview quality",
                owner.get(base + "/chunks").json()["quality_status"] == current["quality_status"],
            )

            spec = importlib.util.spec_from_file_location(
                "m109_rebuild", ROOT / "scripts/rebuild_materials.py"
            )
            rebuild_module = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(rebuild_module)
            candidate = json.loads(
                path.with_suffix(".analysis").joinpath("candidate.json").read_text(encoding="utf-8")
            )
            from final_review.material_conversion import ConvertedMaterial

            converted = ConvertedMaterial(
                markdown=candidate["markdown"],
                sections=candidate["sections"],
                pages=candidate["pages"],
                pipeline="visual-slides-v1",
            )
            isolated_settings = settings.model_copy(
                update={"database_url": SecretStr(isolated_url)}
            )

            def injected(name, *, embedding_failure=False, write_failure=False, conflict=False):
                before = store.get("document", "m109-controlled")
                before_chunks = store.list_material_chunks("m109-controlled")

                class Embeddings(FixedEmbeddings):
                    def embed_documents(self, texts):
                        if embedding_failure:
                            raise RuntimeError("injected embedding failure")
                        return super().embed_documents(texts)

                original_ingest = PostgresStore.ingest

                def ingest(instance, document, rows):
                    original_ingest(instance, document, rows)
                    if write_failure:
                        raise RuntimeError("injected failure after SQL writes")

                def candidate_rebuild(*_args, **_kwargs):
                    if conflict:
                        store.put(
                            "document", "m109-controlled", {**before, "title": "concurrent edit"}
                        )
                    return converted, path.with_suffix(".analysis") / "candidate.json"

                try:
                    with (
                        patch.object(rebuild_module, "Settings", return_value=isolated_settings),
                        patch.object(rebuild_module, "rebuild", side_effect=candidate_rebuild),
                        patch.object(
                            rebuild_module, "OpenAIEmbeddings", side_effect=lambda **_: Embeddings()
                        ),
                        patch.object(PostgresStore, "ingest", ingest),
                        patch.object(
                            sys,
                            "argv",
                            ["rebuild", "--document-id", "m109-controlled", "--publish"],
                        ),
                    ):
                        rebuild_module.main()
                    raise AssertionError("injected failure was accepted")
                except (RuntimeError, ValueError):
                    pass
                after = store.get("document", "m109-controlled")
                expected = {**before, "title": "concurrent edit"} if conflict else before
                check(
                    name,
                    after == expected
                    and store.list_material_chunks("m109-controlled") == before_chunks,
                )
                if conflict:
                    store.put("document", "m109-controlled", before)

            injected("embedding failure retains old version", embedding_failure=True)
            injected("SQL publication failure rolls back all writes", write_failure=True)
            injected("concurrent metadata edit fences publication", conflict=True)
            path2 = directory / "batch-slides.pptx"
            shutil.copy2(path, path2)
            shutil.copytree(path.with_suffix(".analysis"), path2.with_suffix(".analysis"))
            seed("m109-batch", path2)
            cli(
                [
                    "scripts/rebuild_course_materials.py",
                    "--course-id",
                    course["course_id"],
                    "--workers",
                    "2",
                ]
            )
            check(
                "course batch publishes legacy document",
                store.get("document", "m109-batch")["processing_pipeline"] == "visual-slides-v1",
            )
            broken = directory / "broken.pptx"
            broken.write_bytes(b"invalid office file")
            broken_doc, broken_chunks = seed("m109-broken", broken)
            cli(
                [
                    "scripts/rebuild_course_materials.py",
                    "--course-id",
                    course["course_id"],
                    "--workers",
                    "1",
                ],
                expected=1,
            )
            check(
                "course failure keeps old document and chunks",
                store.get("document", "m109-broken") == broken_doc
                and store.list_material_chunks("m109-broken") == broken_chunks,
            )
        report["status"] = "failed" if report.get("failed_visual_checks") else "passed"
    except Exception as exc:
        report["status"] = "failed"
        report["failure_type"] = type(exc).__name__
        raise
    finally:
        if store is not None:
            store.close()
        with psycopg.connect(root_url, autocommit=True) as admin:
            admin.execute(sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(database)))
        report["database_removed"] = True
        report["seconds"] = round(time.monotonic() - started, 2)
        (directory / "report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print("Evidence:", directory, flush=True)
    return 1 if report["status"] == "failed" else 0


if __name__ == "__main__":
    sys.exit(main())
