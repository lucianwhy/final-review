"""Isolated real-PostgreSQL, real-PPT, real-model note-worker benchmark.

Run with ``uv run python scripts/benchmark_note_five_ppt.py``. The script creates
and drops only its own randomly named database. It does not start or restart any
running service. The final JSON contains no credentials or source text.
"""

import json
import time
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import psycopg
from psycopg.conninfo import make_conninfo
from pptx import Presentation
from pptx.util import Inches

from final_review.agent import FinalReviewAgent
from final_review.config import Settings
from final_review.llm import build_models
from final_review.material_conversion import convert_material
from final_review.migrations import apply_migrations
from final_review.note_jobs import process_note_job
from final_review.postgres import PostgresStore
from final_review.rag import KnowledgeBase
from final_review.schemas import AgentRequest, MaterialInput
from final_review.storage import stable_key


def presentation_bytes(file_index: int) -> bytes:
    deck = Presentation()
    deck.slide_width = Inches(13.33)
    deck.slide_height = Inches(7.5)
    topics = [
        "TCP 三次握手确认双向收发能力并协商初始序列号。",
        "HTTP 请求由方法、目标、首部和可选消息体组成。",
        "数据库事务的原子性要求全部提交或全部回滚。",
        "缓存失效策略应同时考虑时效、版本和主动更新。",
        "负载均衡将请求分配到多个实例并监测实例健康。",
    ]
    for slide_index in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(6.5))
        frame = box.text_frame
        frame.clear()
        for paragraph_index in range(12):
            seed = (f"课件 {file_index + 1} 第 {slide_index + 1} 页考点 "
                    f"{paragraph_index + 1}：{topics[file_index]}"
                    "考试作答应先写概念，再说明条件与步骤，并指出常见误区。")
            paragraph = frame.paragraphs[0] if paragraph_index == 0 else frame.add_paragraph()
            paragraph.text = (seed * 11)[:950]
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def main() -> None:
    settings = Settings()
    root_url = settings.database_url.get_secret_value()
    if not root_url:
        raise RuntimeError("DATABASE_URL is required")
    database_name = "final_review_note_bench_" + uuid4().hex[:10]
    benchmark_url = make_conninfo(root_url, dbname=database_name)
    user_id = str(uuid4())
    report = {"database": "isolated PostgreSQL", "selected_files": 5}
    with psycopg.connect(root_url, autocommit=True) as connection:
        connection.execute(f'CREATE DATABASE "{database_name}"')
    store = None
    try:
        apply_migrations(benchmark_url)
        with psycopg.connect(benchmark_url) as connection:
            connection.execute(
                "INSERT INTO app_users(id,email,password_hash) VALUES (%s,%s,%s)",
                (user_id, f"note-bench-{database_name}@example.test", "benchmark"),
            )
        store = PostgresStore(benchmark_url)
        store.setup()
        token = store.bind_user(user_id)
        try:
            store.put("course", "note-bench-course", {
                "course_id": "note-bench-course", "user_id": user_id,
                "name": "五份 PPT 笔记性能验证", "status": "active",
            })
            model, embeddings = build_models(settings, note_generation=True)
            # Parsing the fixed benchmark request need not consume a model call.
            model.note_request = lambda _request: {}
            kb = KnowledgeBase(store, embeddings, settings)
            agent = FinalReviewAgent(store, kb, model, settings)
            documents = []
            ingestion_started = time.monotonic()
            for index in range(5):
                filename = f"benchmark-{index + 1}.pptx"
                content = presentation_bytes(index)
                converted = convert_material(content, filename,
                                             settings.max_upload_mb * 1024 * 1024)
                document_id = f"benchmark-document-{index + 1}"
                prepared, chunks = kb.prepare(MaterialInput(
                    document_id=document_id, course_id="note-bench-course",
                    title=f"课件 {index + 1}", source_type="teacher_ppt",
                    markdown=converted.markdown,
                ), source_origin="user_upload", sections=converted.sections)
                store.ingest({**prepared, "user_id": user_id, "parse_status": "ready",
                              "file_name": filename}, chunks)
                documents.append(document_id)
                report.setdefault("file_chunks", []).append(len(chunks))
            report["ppt_conversion_and_index_seconds"] = round(
                time.monotonic() - ingestion_started, 2)
            initial = agent.invoke(AgentRequest(
                course_id="note-bench-course", session_id="benchmark-session",
                message="生成笔记", intent="note",
            ), user_id)
            if initial.status != "needs_input":
                raise RuntimeError(f"Expected needs_input, got {initial.status}")
            now = datetime.now(UTC).isoformat()
            job_id = stable_key("note-bench-course", "benchmark-session", "benchmark-job")
            store.put("note_job", job_id, {
                "job_id": job_id, "user_id": user_id, "course_id": "note-bench-course",
                "conversation_id": "benchmark-conversation", "session_id": "benchmark-session",
                "event_id": "benchmark-event", "note_input": {
                    "note_type": "key_points", "scope": "", "duration_minutes": 10,
                    "source_document_ids": documents,
                }, "note_policy_version": 2, "status": "queued", "attempts": 0,
                "max_attempts": 3, "available_at": now, "lease_until": None,
                "created_at": now,
            })
            generation_calls = 0
            verification_calls = 0
            original_note, original_verify = model.note, model.verify

            def counted_note(data):
                nonlocal generation_calls
                generation_calls += 1
                return original_note(data)

            def counted_verify(data):
                nonlocal verification_calls
                verification_calls += 1
                return original_verify(data)

            model.note, model.verify = counted_note, counted_verify
            started = time.monotonic()
            claimed = store.claim_note_job()
            if not claimed or claimed["job_id"] != job_id:
                raise RuntimeError("Could not claim benchmark note job")
            process_note_job(store, agent, claimed)
            report["note_worker_seconds"] = round(time.monotonic() - started, 2)
            job = store.get("note_job", job_id)
            report.update({
                "status": job["status"], "error": job.get("error"),
                "generation_calls": generation_calls,
                "verification_calls": verification_calls,
                "completed_batches": len(job.get("partial_batches", {})),
                "fallback_batches": sum(bool(batch.get("fallback_reason"))
                                        for batch in job.get("partial_batches", {}).values()),
                "coverage": (job.get("selection_plan") or {}).get("coverage"),
            })
            revisions = store.scan("asset_revision", {"course_id": "note-bench-course"})
            report["draft_points"] = len(revisions[0]["points"]) if revisions else 0
        finally:
            store.reset_user(token)
    finally:
        if store:
            store.close()
        with psycopg.connect(root_url, autocommit=True) as connection:
            connection.execute(f'DROP DATABASE "{database_name}" WITH (FORCE)')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
